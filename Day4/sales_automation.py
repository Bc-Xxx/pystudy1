# ================== Day 4 完整小项目：销售数据自动化处理工具 ==================
import pandas as pd
import pathlib
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime


# ================== 1. 美化函数（复用） ==================
def 美化_excel(文件名):
    """美化 Excel 表格标题行"""
    wb = load_workbook(文件名)
    ws = wb.active

    title_font = Font(bold=True, color="FFFFFF", size=13)
    title_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")

    for cell in ws[1]:
        cell.font = title_font
        cell.fill = title_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # 自动调整列宽
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        ws.column_dimensions[column_letter].width = max_length + 4

    wb.save(文件名)
    print(f"🎨 已美化: {文件名}")


# ================== 2. 核心处理函数 ==================
def 处理销售数据(输入文件):
    """处理单个销售数据文件"""
    print(f"\n正在处理: {输入文件}")

    # 读取
    df = pd.read_excel(输入文件)

    # 处理
    df['总价'] = df['销量'] * df['单价']
    df = df[df['销量'] > 50]  # 筛选
    df = df.sort_values(by='总价', ascending=False)  # 排序

    # 生成输出文件名
    输出文件名 = f"处理后_{输入文件.name}"
    输出路径 = 输入文件.parent / 输出文件名

    # 保存
    df.to_excel(输出路径, index=False)

    # 美化
    美化_excel(str(输出路径))

    print(f"✅ 处理完成: {输出文件名}")
    return df


# ================== 3. 批量处理 ==================
def 批量处理文件夹(文件夹路径='.'):
    folder = pathlib.Path(文件夹路径)
    excel_files = list(folder.glob('*.xlsx'))

    print(f"共发现 {len(excel_files)} 个 Excel 文件\n")

    for file in excel_files:
        if file.name.startswith('~$') or file.name.startswith('处理后_'):
            continue
        处理销售数据(file)


# ================== 运行主程序 ==================
if __name__ == "__main__":
    print("🚀 销售数据自动化处理工具启动！")
    print("=" * 50)

    批量处理文件夹('.')  # 处理当前文件夹所有文件

    print("\n🎉 所有文件处理完成！")
    print(f"完成时间: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
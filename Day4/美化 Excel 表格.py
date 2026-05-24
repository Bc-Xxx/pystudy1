import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

def 美化_excel(文件名):
    wb=load_workbook(文件名)
    ws=wb.active

    # 设置标题行样式
    title_font=Font(bold=True,color='FFFFFF',size=12)
    title_fill=PatternFill(start_color="366092", end_color="366092", fill_type="solid")

    # 遍历第一行（标题行）
    for cell in ws[1]:
            cell.font = title_font
            cell.fill = title_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')

    # 自动调整列宽
    for column in ws.columns:
        max_length=0
        column_letter=column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        ws.column_dimensions[column_letter].width = adjusted_width

    wb.save(文件名)  # 保存修改
    print(f"✅ 美化完成！文件：{文件名}")

# 使用美化函数
美化_excel('Day4_Excel/销售数据.xlsx')
print("美化后的文件已保存！")